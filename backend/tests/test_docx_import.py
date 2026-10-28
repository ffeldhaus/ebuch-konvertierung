"""
Tests for the DOCX → Markdown import used by the Feedback and Benchmark modes.

Background: Manually corrected transfers are usually uploaded as Word documents. If the
import drops heading levels or list nesting, the Gemma feedback analysis compares against a
distorted correction and produces wrong diagnoses (e.g. "task labels should be plain text"
while they actually were level-5 headings). All documents here are synthetic.
"""
from docx import Document
from docx.shared import Cm

from backend.services.docx_service import extract_markdown_from_docx


def _build_docx(path, paragraphs):
    """Creates a DOCX from (style_name, text, left_indent_cm|None) tuples using built-in styles."""
    doc = Document()
    for style, text, indent in paragraphs:
        p = doc.add_paragraph(text, style=style)
        if indent is not None:
            p.paragraph_format.left_indent = Cm(indent)
    doc.save(path)
    return path


def test_docx_import_keeps_heading_levels_5_and_6(tmp_path):
    path = _build_docx(str(tmp_path / "h.docx"), [
        ("Heading 4", "Kompetenzbereich Beispiel", None),
        ("Heading 5", "X1", None),
        ("Normal", "Beschreiben Sie das Experiment.", None),
        ("Heading 6", "Detailebene", None),
    ])
    md = extract_markdown_from_docx(path)
    lines = md.split("\n\n")
    assert "#### Kompetenzbereich Beispiel" in lines
    assert "##### X1" in lines
    assert "###### Detailebene" in lines
    assert "Beschreiben Sie das Experiment." in lines


def test_docx_import_nested_list_styles_are_indented(tmp_path):
    """'List 2' / 'List Bullet 2' are second-level items → indented by two spaces (export convention)."""
    path = _build_docx(str(tmp_path / "l.docx"), [
        ("List", "-\tOberbegriff", None),
        ("List 2", "-\tUnterbegriff", None),
        ("List Bullet 2", "Noch ein Unterbegriff", None),
        ("List", "a)\tTeilaufgabe", None),
    ])
    md = extract_markdown_from_docx(path)
    lines = md.split("\n\n")
    assert "-\tOberbegriff" in lines
    assert "  -\tUnterbegriff" in lines
    assert "  -\tNoch ein Unterbegriff" in lines
    # Plain 'List' style items that already carry their own marker stay untouched
    assert "a)\tTeilaufgabe" in lines


def test_docx_import_list_level_from_left_indent(tmp_path):
    """Our own export marks nesting only via left indent (1 cm per level) on style 'List'."""
    path = _build_docx(str(tmp_path / "i.docx"), [
        ("List", "-\tEbene eins", 1.0),
        ("List", "-\tEbene zwei", 2.0),
    ])
    md = extract_markdown_from_docx(path)
    lines = md.split("\n\n")
    assert "-\tEbene eins" in lines
    assert "  -\tEbene zwei" in lines


def test_docx_roundtrip_heading5_and_nested_list(tmp_path):
    from backend.services.docx_service import convert_markdown_to_docx

    src = "#### Bereich\n\n##### X2\n\nText der Aufgabe.\n\n-\tPunkt\n  -\tUnterpunkt"
    out = str(tmp_path / "rt.docx")
    convert_markdown_to_docx(src, out)
    md = extract_markdown_from_docx(out)
    lines = md.split("\n\n")
    assert "##### X2" in lines
    assert "-\tPunkt" in lines
    assert "  -\tUnterpunkt" in lines


def test_docx_roundtrip_multiline_bild_block(tmp_path):
    """The multi-line <Bild> block is one Word paragraph with soft line breaks (JAWS: Strg+Pfeil skips it)."""
    from backend.services.docx_service import convert_markdown_to_docx

    block = "<Bild> Strukturdiagramm:\nBeschreibung mit Bildtext: Start \\to Ziel\n</Bild>"
    out = str(tmp_path / "bild.docx")
    convert_markdown_to_docx(f"# Titel\n\n{block}\n\nText.", out)
    doc = Document(out)
    bild_paras = [p for p in doc.paragraphs if p.text.startswith("<Bild>")]
    assert len(bild_paras) == 1
    assert bild_paras[0].text == block
    assert block in extract_markdown_from_docx(out)
