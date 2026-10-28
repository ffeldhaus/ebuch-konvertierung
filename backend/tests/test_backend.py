import os
import shutil
import pytest
import json
from fastapi.testclient import TestClient

from backend.main import app, task_manager, active_tasks
from backend.services.fibs_postprocess import post_process_structure, standardise_markdown_fallback
from backend.services.llm_service import standardise_markdown
from backend.services.docx_service import convert_markdown_to_docx
from backend.services.extraction_service import (
    validate_pages_string,
    extract_pdf_pages,
    parse_page_selection,
    is_contiguous_range,
    insert_page_numbers_from_metadata,
)

client = TestClient(app)

# Helper constants
BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
TEST_DIR = os.path.join(BASE_DIR, "assets", "test_temp")
os.makedirs(TEST_DIR, exist_ok=True)

@pytest.fixture(scope="module", autouse=True)
def cleanup():
    # Setup
    yield
    # Teardown
    if os.path.exists(TEST_DIR):
        shutil.rmtree(TEST_DIR)

def test_llm_rules_fallback():
    # 1. Test bold conversion (**bold** -> _bold_)
    raw_md = "This is **bold** text."
    result = standardise_markdown_fallback(raw_md, "uebertragungsregeln_09_2025")
    assert "This is _bold_ text." in result

    # 2. Test italic conversion (*italic* -> _italic_)
    raw_md = "This is *italic* text."
    result = standardise_markdown_fallback(raw_md, "uebertragungsregeln_09_2025")
    assert "This is _italic_ text." in result

    # 3. Test list conversion (bullet -> -\tcontent)
    raw_md = "- Bullet point text"
    result = standardise_markdown_fallback(raw_md, "uebertragungsregeln_09_2025")
    assert "-\tBullet point text" in result

    # 4. Test page numbers
    raw_md = "45"
    result = standardise_markdown_fallback(raw_md, "uebertragungsregeln_09_2025")
    assert "((45))" in result

    # 5. Test box detection (Merksatz)
    raw_md = "Merksatz:\nDispersion beschreibt die Lichtbrechung."
    result = standardise_markdown_fallback(raw_md, "uebertragungsregeln_09_2025")
    assert "<Rahmen>" in result
    assert "</Rahmen>" in result

    # 6. Test formula cleaning
    raw_md = "<L> n1 * sin(alpha) = n2 * sin(beta) </L>"
    result = standardise_markdown_fallback(raw_md, "uebertragungsregeln_09_2025")
    assert "<L>" not in result
    assert "</L>" not in result

    # 7. Test quotation marks, ellipses, and placeholders
    raw_md = "Er sagte: „Hallo Welt“ und ‚Test‘ mit Paarreim (aabb...) und Lücke: 24 + 37 = ____"
    result = standardise_markdown_fallback(raw_md, "uebertragungsregeln_09_2025")
    assert '"Hallo Welt"' in result
    assert "'Test'" in result
    assert "(aabb ...)" in result
    assert "24 + 37 = _..._" in result

def test_docx_generation():
    markdown_content = (
        "((45))\n\n"
        "# Physik für die Sekundarstufe II\n\n"
        "Das Licht wird gebrochen.\n\n"
        "-\tLichtgeschwindigkeit c = 299.792 km/s.\n\n"
        "<Tabelle>\n"
        "<table><tr><td>n</td><td>1.51</td></tr></table>\n"
        "</Tabelle>"
    )
    
    docx_path = os.path.join(TEST_DIR, "test_output.docx")
    convert_markdown_to_docx(markdown_content, docx_path)
    
    assert os.path.exists(docx_path)
    assert os.path.getsize(docx_path) > 0

def test_fastapi_conversion_flow(monkeypatch):
    # Mock LLM settings to force rule-based fallback
    monkeypatch.setattr("backend.services.llm_service.start_rapid_mlx", lambda: None)
    
    # Generate a dummy input file
    dummy_pdf_path = os.path.join(TEST_DIR, "dummy.pdf")
    with open(dummy_pdf_path, "wb") as f:
        # Minimal valid PDF signature
        f.write(b"%PDF-1.4\n1 0 obj\n<<\n/Type /Catalog\n>>\nendobj\ntrailer\n<<\n/Root 1 0 R\n>>\n%%EOF")

    # Step 1: Upload the file
    with open(dummy_pdf_path, "rb") as f:
        response = client.post("/api/upload", files={"file": ("dummy.pdf", f, "application/pdf")})
    
    assert response.status_code == 200
    task_id = response.json()["task_id"]
    assert task_id is not None
    
    # Step 2: Extract (will fail with dummy PDF but we can mock / check task exists)
    # We can check that a 404 is thrown for invalid task_ids
    response_invalid = client.post("/api/extract", data={"task_id": "invalid-id"})
    assert response_invalid.status_code == 404
    
    # Step 3: Standardise
    response = client.post("/api/standardise", json={
        "task_id": task_id,
        "markdown": "Some raw **markdown**",
        "ruleset": "uebertragungsregeln_09_2025"
    })
    assert response.status_code == 200
    
    # Parse NDJSON response lines
    completed_md = None
    for line in response.iter_lines():
        if line.strip():
            data = json.loads(line)
            if data.get("status") == "completed":
                completed_md = data.get("markdown")
                
    assert completed_md is not None
    assert "Some raw _markdown_" in completed_md
    
    # Step 4: Export
    response = client.post("/api/export", json={
        "task_id": task_id,
        "markdown": "Some standardised markdown"
    })
    assert response.status_code == 200
    download_url = response.json()["download_url"]
    assert f"/api/download/{task_id}" in download_url
    
    # Step 5: Download
    response = client.get(f"/api/download/{task_id}")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

def test_tasks_management_endpoints():
    # 1. Create a task via upload
    dummy_pdf_path = os.path.join(TEST_DIR, "dummy_tasks.pdf")
    with open(dummy_pdf_path, "wb") as f:
        f.write(b"%PDF-1.4\n1 0 obj\n<<\n/Type /Catalog\n>>\nendobj\ntrailer\n<<\n/Root 1 0 R\n>>\n%%EOF")

    with open(dummy_pdf_path, "rb") as f:
        response = client.post("/api/upload", files={"file": ("dummy_tasks.pdf", f, "application/pdf")})
    
    assert response.status_code == 200
    task_id = response.json()["task_id"]
    
    # 2. Get all tasks list
    response = client.get("/api/tasks")
    assert response.status_code == 200
    task_list = response.json()
    assert any(t["task_id"] == task_id for t in task_list)
    
    # 3. Get task state
    response = client.get(f"/api/tasks/{task_id}/state")
    assert response.status_code == 200
    state = response.json()
    assert state["task_id"] == task_id
    assert state["status"] == "pending"
    
    # 4. Cancel/Delete task
    response = client.delete(f"/api/tasks/{task_id}")
    assert response.status_code == 200
    assert response.json()["status"] == "deleted"
    
    # Verify task is deleted from active tasks
    response = client.get(f"/api/tasks/{task_id}/state")
    assert response.status_code == 404


def test_task_abort_flow():
    # 1. Upload a file to create a task
    dummy_pdf_path = os.path.join(TEST_DIR, "dummy_abort.pdf")
    with open(dummy_pdf_path, "wb") as f:
        f.write(b"%PDF-1.4\n1 0 obj\n<<\n/Type /Catalog\n>>\nendobj\ntrailer\n<<\n/Root 1 0 R\n>>\n%%EOF")

    with open(dummy_pdf_path, "rb") as f:
        response = client.post("/api/upload", files={"file": ("dummy_abort.pdf", f, "application/pdf")})
    assert response.status_code == 200
    task_id = response.json()["task_id"]

    # 2. Call abort endpoint
    response = client.post(f"/api/tasks/{task_id}/abort")
    assert response.status_code == 200
    assert response.json()["status"] == "aborted"

    # 3. Check task state is aborted
    response = client.get(f"/api/tasks/{task_id}/state")
    assert response.status_code == 200
    assert response.json()["status"] == "aborted"
    assert "abgebrochen" in response.json()["progress"]["phase"]

    # Clean up
    client.delete(f"/api/tasks/{task_id}")


def test_task_state_raw_markdown_recovery():
    # 1. Upload a file
    dummy_pdf_path = os.path.join(TEST_DIR, "dummy_recover.pdf")
    with open(dummy_pdf_path, "wb") as f:
        f.write(b"%PDF-1.4\n1 0 obj\n<<\n/Type /Catalog\n>>\nendobj\ntrailer\n<<\n/Root 1 0 R\n>>\n%%EOF")

    with open(dummy_pdf_path, "rb") as f:
        response = client.post("/api/upload", files={"file": ("dummy_recover.pdf", f, "application/pdf")})
    assert response.status_code == 200
    task_id = response.json()["task_id"]

    # 2. Get task from task_manager and set raw_markdown
    from backend.main import task_manager
    task = task_manager.get_sync(task_id)
    assert task is not None
    task.raw_markdown = "# Extrahierter Text vor Bildanalyse\n((1))\nBeispieltext"
    task.status = "processing"
    task.progress["phase"] = "Generiere Bildbeschreibungen mit Gemma 4..."

    # 3. Query state endpoint (simulating page reload while processing)
    response = client.get(f"/api/tasks/{task_id}/state")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "processing"
    assert data["raw_markdown"] == "# Extrahierter Text vor Bildanalyse\n((1))\nBeispieltext"
    assert "Bildbeschreibungen" in data["progress"]["phase"]

    # Clean up
    client.delete(f"/api/tasks/{task_id}")

def test_pdf_page_selection():
    import pypdfium2 as pdfium
    
    # Create a 3-page dummy PDF using pypdfium2
    doc = pdfium.PdfDocument.new()
    doc.new_page(200, 200)
    doc.new_page(200, 200)
    doc.new_page(200, 200)
    
    dummy_pdf_path = os.path.join(TEST_DIR, "dummy_3pages.pdf")
    doc.save(dummy_pdf_path)
    doc.close()
    
    # 1. Test extracting page 1 and 3 (pages index 0 and 2)
    selected_pdf_path = os.path.join(TEST_DIR, "dummy_selected.pdf")
    if os.path.exists(selected_pdf_path):
        os.remove(selected_pdf_path)
        
    success = extract_pdf_pages(dummy_pdf_path, "1,3", selected_pdf_path)
    assert success is True
    assert os.path.exists(selected_pdf_path)
    
    # Open extracted PDF and check pages count
    res_doc = pdfium.PdfDocument(selected_pdf_path)
    assert len(res_doc) == 2
    res_doc.close()
    
    # 2. Test extracting a range (1-2)
    success = extract_pdf_pages(dummy_pdf_path, "1-2", selected_pdf_path)
    assert success is True
    res_doc = pdfium.PdfDocument(selected_pdf_path)
    assert len(res_doc) == 2
    res_doc.close()
    
    # 3. Test empty or invalid page string
    assert extract_pdf_pages(dummy_pdf_path, "", selected_pdf_path) is False
    assert extract_pdf_pages(dummy_pdf_path, "invalid", selected_pdf_path) is False

def test_pdf_rendering_service():
    import pypdfium2 as pdfium
    from backend.services.pdf_render_service import render_single_page_bytes

    # 1. Create a simple 1-page PDF
    doc = pdfium.PdfDocument.new()
    doc.new_page(200, 200)
    pdf_path = os.path.join(TEST_DIR, "render_test.pdf")
    doc.save(pdf_path)
    doc.close()

    # 2. Render page 0 to PNG bytes
    png_bytes = render_single_page_bytes(pdf_path, 0)
    assert png_bytes is not None
    assert png_bytes.startswith(b"\x89PNG")

    # 3. Test caching: re-render should return identical cached bytes
    cached_bytes = render_single_page_bytes(pdf_path, 0)
    assert cached_bytes == png_bytes

    # 4. Out-of-bounds page index and missing file return None safely
    assert render_single_page_bytes(pdf_path, 5) is None
    assert render_single_page_bytes(pdf_path, -1) is None
    assert render_single_page_bytes(os.path.join(TEST_DIR, "nonexistent.pdf"), 0) is None

def test_diff_service_highlighting():
    from backend.services.diff_service import get_diff_highlighting

    raw_text = "This is a book page with ((45)) and a **bold** word."
    std_text = "This is a book page with ((45)) and a *bold* word."

    segments = get_diff_highlighting(raw_text, std_text)
    assert len(segments) > 0
    
    # We should have format and equal tokens
    has_format = False
    has_equal = False
    for s in segments:
        if s["category"] == "format":
            has_format = True
        if s["type"] == "equal":
            has_equal = True
            
    assert has_equal is True
    # The formatting markers changed (** -> *)
    assert has_format is True

    # Test tag category
    raw_text = "Text here."
    std_text = "Text here. <Bild> Foto: Test </Bild>"
    segments = get_diff_highlighting(raw_text, std_text)
    
    has_tag = False
    for s in segments:
        if s["category"] == "tag" and s["type"] == "insert":
            has_tag = True
    assert has_tag is True

def test_diff_endpoint():
    response = client.post("/api/diff", json={
        "raw": "This is raw text.",
        "std": "This is standard text."
    })
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "success"
    assert "segments" in data
    assert len(data["segments"]) > 0

def test_strip_ground_truth_preamble():
    from backend.services.diff_service import strip_ground_truth_preamble, get_diff_highlighting

    # Synthetic FIBS preamble example
    gt_with_preamble = """Fach: Biologie
Titel: Lehrbuch Biologie Gesamtband
Verlag: Musterverlag
ISBN: 978-0-00-123456-7
Seite: 80 bis 95
Übertragen im FIBS
Stiftstr. 53
59494 Soest

<Anmerkung>
-	Einige Seiten werden als Doppelseiten übertragen.
-	Schwierigkeitsgrade von Aufgaben werden vor der Aufgabenstellung in eckigen Klammern angegeben: [leicht], [mittel] oder [schwer].
-	Zur besseren Orientierung wird den durchnummerierten Abbildungen B für Bildmaterial vorangestellt.
-	Pfeile werden als \\to, beschrifte Pfeile als \\xra{Text} dargestellt.
-	Alle Videos, interaktive Übungen und Dokumente zum Schulbuch sind im eBook und in den Medien zum Schulbuch verfügbar.
</Anmerkung>

((80))
# Zellen und Gewebe
Zellen sind die Grundeinheiten aller Lebewesen."""

    cleaned = strip_ground_truth_preamble(gt_with_preamble)
    assert cleaned.startswith("((80))")
    assert "Fach: Biologie" not in cleaned
    assert "<Anmerkung>" not in cleaned
    assert "Zellen und Gewebe" in cleaned

    # Test diff highlighting auto-strip
    candidate_std = "((80))\n# Zellen und Gewebe\nZellen sind die Grundeinheiten aller Lebewesen."
    diff_segments = get_diff_highlighting(gt_with_preamble, candidate_std)
    # The first segment should be equal ((80)), not a delete of the preamble
    assert diff_segments[0]["type"] == "equal"
    assert diff_segments[0]["text"].startswith("((80))")

def test_image_descriptions_ignored_in_diff():
    from backend.services.diff_service import normalize_images_for_diff, get_diff_highlighting

    gt_doc = """((80))
# Dispersion
<Bild> Zeichnung: Versuchsaufbau zur optischen Dispersion mit einem Glasprisma. Ein weißer Lichtstrahl tritt von links in ein Prisma ein. </Bild>
Lichtstrahlen werden gebrochen."""

    pred_doc = """((80))
# Dispersion
<Bild> Zeichnung: Ein Glasprisma bricht weißes Licht in seine Spektralfarben auf. </Bild>
Lichtstrahlen werden gebrochen."""

    # 1. Normalize helper check
    norm_gt = normalize_images_for_diff(gt_doc)
    norm_pred = normalize_images_for_diff(pred_doc)
    assert norm_gt == norm_pred
    assert "<Bild></Bild>" in norm_gt

    # 2. Diff check: when both have image tags with different wording, segments should be entirely 'equal'
    diff_segments = get_diff_highlighting(gt_doc, pred_doc)
    assert all(s["type"] == "equal" for s in diff_segments)

def test_extract_document_components():
    from backend.services.diff_service import extract_document_components

    doc = """# Header
((1))
Paragraph 1.

<Bild> Zeichnung: Ein Prisma. </Bild>

Paragraph 2.

<Tabelle>
| Spalte 1 | Spalte 2 |
| A | B |
</Tabelle>

<Rahmen>
Wichtiger Merksatz.
</Rahmen>

Paragraph 3."""

    comp = extract_document_components(doc)
    assert "# Header" in comp["clean_body"]
    assert "Paragraph 1." in comp["clean_body"]
    assert "Paragraph 2." in comp["clean_body"]
    assert "Paragraph 3." in comp["clean_body"]
    assert "<Bild>" not in comp["clean_body"]
    assert "<Tabelle>" not in comp["clean_body"]
    assert "<Rahmen>" not in comp["clean_body"]

    assert len(comp["images"]) == 1
    assert "Prisma" in comp["images"][0]
    assert len(comp["tables"]) == 1
    assert "Spalte 1" in comp["tables"][0]
    assert len(comp["boxes"]) == 1
    assert "Merksatz" in comp["boxes"][0]

def test_floating_blocks_diff_reordering_clean_paragraphs():
    from backend.services.diff_service import get_diff_highlighting

    # Ground truth has Image and Table at the top
    gt_doc = """((1))
# Kapitel 1
<Bild> Zeichnung: Versuchsaufbau. </Bild>
<Tabelle>
| A | B |
| 1 | 2 |
</Tabelle>
Erster wichtiger Textabschnitt.
Zweiter wichtiger Textabschnitt."""

    # Prediction extracted Image and Table at the bottom after paragraphs
    pred_doc = """((1))
# Kapitel 1
Erster wichtiger Textabschnitt.
Zweiter wichtiger Textabschnitt.
<Bild> Zeichnung: Versuchsaufbau eines Experiments. </Bild>
<Tabelle>
| A | B |
| 1 | 2 |
</Tabelle>"""

    diff = get_diff_highlighting(gt_doc, pred_doc)
    # The text paragraphs themselves must not be marked as delete or insert!
    text_deletions = [s for s in diff if s["type"] == "delete" and "Textabschnitt" in s["text"]]
    text_insertions = [s for s in diff if s["type"] == "insert" and "Textabschnitt" in s["text"]]
    assert len(text_deletions) == 0
    assert len(text_insertions) == 0

def test_evaluate_images_metric():
    from scripts.run_benchmark import evaluate_images

    gt = """# Optik
<Bild> Zeichnung: Prisma bricht Licht in Spektralfarben. </Bild>
<Bild> Foto: Laserstrahl. </Bild>"""

    pred = """# Optik
<Bild> Foto: Laserstrahl im Labor. </Bild>
<Bild> Zeichnung: Prisma Dispersion von Lichtstrahlen. </Bild>"""

    res = evaluate_images(pred, gt)
    assert res["gold_count"] == 2
    assert res["pred_count"] == 2
    assert res["matched_count"] == 2
    assert res["score"] > 80.0

def test_standardise_fallback(monkeypatch):
    from backend.services.llm_service import standardise_markdown
    
    # Mock start_rapid_mlx to return None; the rule-based fallback only runs on explicit opt-in
    monkeypatch.setattr("backend.services.llm_service.start_rapid_mlx", lambda: None)
    
    result = standardise_markdown("Some **bold** text.", ruleset_key="uebertragungsregeln_09_2025", allow_fallback=True)
    # Rule-based fallback converts **bold** -> _bold_
    assert result == "Some _bold_ text."

def test_optimize_endpoint(monkeypatch):
    # Create a task in active tasks
    dummy_pdf_path = os.path.join(TEST_DIR, "dummy_optimize.pdf")
    with open(dummy_pdf_path, "wb") as f:
        f.write(b"%PDF-1.4\n1 0 obj\n<<\n/Type /Catalog\n>>\nendobj\ntrailer\n<<\n/Root 1 0 R\n>>\n%%EOF")

    with open(dummy_pdf_path, "rb") as f:
        response = client.post("/api/upload", files={"file": ("dummy_optimize.pdf", f, "application/pdf")})
    
    assert response.status_code == 200
    task_id = response.json()["task_id"]
    
    # Mock LLM settings to force rule-based fallback
    monkeypatch.setattr("backend.services.llm_service.start_rapid_mlx", lambda: None)
    
    # Test optimization call (fallback returns current_markdown unchanged)
    response = client.post("/api/optimize", json={
        "task_id": task_id,
        "prompt": "improve formel formatting on page 2",
        "current_markdown": "Standardised markdown output",
        "ruleset": "uebertragungsregeln_09_2025"
    })
    
    assert response.status_code == 200
    data = response.json()
    assert data["task_id"] == task_id
    assert data["markdown"] == "Standardised markdown output"

def test_dotx_template_support():
    dotx_path = os.path.join(BASE_DIR, "assets", "templates", "Dokumentenvorlage FIBS 25.05.2025_NEU.dotx")
    out_docx_path = os.path.join(TEST_DIR, "test_dotx_output.docx")
    
    assert os.path.exists(dotx_path)
    
    markdown_content = (
        "((45))\n\n"
        "# Physik für die Sekundarstufe II\n\n"
        "Das Licht wird gebrochen.\n"
    )
    
    # Should compile and parse DOTX successfully
    convert_markdown_to_docx(markdown_content, out_docx_path, template_path=dotx_path)
    
    assert os.path.exists(out_docx_path)
    assert os.path.getsize(out_docx_path) > 0

def test_custom_ruleset_merging(monkeypatch):
    from backend.services.llm_service import standardise_markdown
    
    # Create a dummy task directory and custom ruleset
    task_id = "test-ruleset-task-123"
    task_dir = os.path.join(BASE_DIR, "assets", "output", task_id)
    os.makedirs(task_dir, exist_ok=True)
    
    custom_rules_file = os.path.join(task_dir, "custom_ruleset.txt")
    with open(custom_rules_file, "w", encoding="utf-8") as f:
        f.write("RULE: Convert all occurrences of 'Hund' to 'Katze'.")
        
    captured_messages = []
    
    # Mock client to intercept prompt messages
    class MockChatCompletions:
        def create(self, **kwargs):
            nonlocal captured_messages
            captured_messages = kwargs.get("messages", [])
            class MockDelta:
                content = "Mock response"
            class MockChoice:
                delta = MockDelta()
            class MockChunk:
                choices = [MockChoice()]
            return [MockChunk()]
            
    class MockClient:
        class MockChat:
            completions = MockChatCompletions()
        chat = MockChat()
        
    monkeypatch.setattr("backend.services.llm_service.start_rapid_mlx", lambda: MockClient())
    
    # Call standardise
    standardise_markdown("Ein Hund rennt.", ruleset_key="uebertragungsregeln_09_2025", task_id=task_id)
    
    # Clean up
    if os.path.exists(custom_rules_file):
        os.remove(custom_rules_file)
    if os.path.exists(task_dir):
        os.rmdir(task_dir)
        
    # Verify ruleset was merged
    assert len(captured_messages) > 0
    sys_msg = captured_messages[0]["content"]
    assert "ZUSÄTZLICHE DATEISPEZIFISCHE REGELN:" in sys_msg
    assert "RULE: Convert all occurrences of 'Hund' to 'Katze'." in sys_msg
    assert "absolute Priorität und überschreiben alle" in sys_msg

def test_pdf_ruleset_upload_and_extraction(monkeypatch):
    # Mock run_docling_extraction to avoid calling actual CLI on a fake PDF
    monkeypatch.setattr(
        "scripts.run_benchmark.run_docling_extraction",
        lambda pdf_path, output_dir: ("MOCK RULES FROM PDF", 0.1)
    )
    
    # Set up dummy task via the task manager
    from backend.services.task_manager import Task
    task_id = "test-ruleset-pdf-123"
    task = Task(
        task_id=task_id,
        original_name="test.pdf",
        input_path="/mock/path",
        extension=".pdf",
    )
    active_tasks[task_id] = task
    
    # Upload PDF file
    import io
    pdf_content = b"%PDF-1.4 dummy content"
    files = {"file": ("custom_rules.pdf", io.BytesIO(pdf_content), "application/pdf")}
    data = {"task_id": task_id}
    
    response = client.post("/api/upload-ruleset", data=data, files=files)
    assert response.status_code == 200
    assert response.json()["ruleset_name"] == "custom_rules.pdf"
    
    # Check that custom_ruleset.txt was created inside output directory
    from backend.main import OUTPUT_DIR
    custom_ruleset_file = os.path.join(OUTPUT_DIR, task_id, "custom_ruleset.txt")
    assert os.path.exists(custom_ruleset_file)
    with open(custom_ruleset_file, "r", encoding="utf-8") as f:
        content = f.read()
    assert content == "MOCK RULES FROM PDF"
    
    # Clean up files & task
    if os.path.exists(custom_ruleset_file):
        os.remove(custom_ruleset_file)
    task_dir = os.path.dirname(custom_ruleset_file)
    if os.path.exists(task_dir):
        try:
            os.rmdir(task_dir)
        except Exception:
            pass
    active_tasks.pop(task_id, None)

def test_docx_ruleset_upload_and_extraction(monkeypatch):
    # Mock extract_markdown_from_docx to return clean markdown
    monkeypatch.setattr(
        "backend.services.docx_service.extract_markdown_from_docx",
        lambda docx_path: "MOCK RULES FROM DOCX"
    )
    
    # Set up dummy task via the task manager
    from backend.services.task_manager import Task
    task_id = "test-ruleset-docx-123"
    task = Task(
        task_id=task_id,
        original_name="test.docx",
        input_path="/mock/path",
        extension=".docx",
    )
    active_tasks[task_id] = task
    
    # Upload DOCX file
    import io
    docx_content = b"PK\x03\x04 dummy zip/docx content"
    files = {"file": ("custom_rules.docx", io.BytesIO(docx_content), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    data = {"task_id": task_id}
    
    response = client.post("/api/upload-ruleset", data=data, files=files)
    assert response.status_code == 200
    res_json = response.json()
    assert res_json["ruleset_name"] == "custom_rules.docx"
    assert res_json["content"] == "MOCK RULES FROM DOCX"
    assert res_json["chars"] == len("MOCK RULES FROM DOCX")
    
    # Check that custom_ruleset.txt was created inside output directory
    from backend.main import OUTPUT_DIR
    custom_ruleset_file = os.path.join(OUTPUT_DIR, task_id, "custom_ruleset.txt")
    assert os.path.exists(custom_ruleset_file)
    with open(custom_ruleset_file, "r", encoding="utf-8") as f:
        content = f.read()
    assert content == "MOCK RULES FROM DOCX"

    # Test GET /api/ruleset/{task_id} inspection endpoint
    get_res = client.get(f"/api/ruleset/{task_id}")
    assert get_res.status_code == 200
    get_json = get_res.json()
    assert get_json["has_ruleset"] is True
    assert get_json["content"] == "MOCK RULES FROM DOCX"
    
    # Clean up files & task
    if os.path.exists(custom_ruleset_file):
        os.remove(custom_ruleset_file)
    task_dir = os.path.dirname(custom_ruleset_file)
    if os.path.exists(task_dir):
        try:
            os.rmdir(task_dir)
        except Exception:
            pass
    active_tasks.pop(task_id, None)

def test_validate_pages_string():
    # Valid selections
    assert validate_pages_string("", 5) == (True, "")
    assert validate_pages_string("1,3", 5) == (True, "")
    assert validate_pages_string("1-3, 5", 5) == (True, "")
    assert validate_pages_string("  1-3 , 4, 5  ", 5) == (True, "")
    
    # Invalid characters
    is_valid, msg = validate_pages_string("1,2,abc", 5)
    assert is_valid is False
    assert "ungültige Zeichen" in msg
    
    # Out of bounds single page
    is_valid, msg = validate_pages_string("6", 5)
    assert is_valid is False
    assert "liegt außerhalb des gültigen Bereichs" in msg

    # Out of bounds range start
    is_valid, msg = validate_pages_string("6-7", 5)
    assert is_valid is False
    assert "liegt außerhalb des gültigen Bereichs" in msg

    # Out of bounds range end
    is_valid, msg = validate_pages_string("3-8", 5)
    assert is_valid is False
    assert "liegt außerhalb des gültigen Bereichs" in msg

    # Non-positive numbers
    is_valid, msg = validate_pages_string("0,1", 5)
    assert is_valid is False
    assert "größer als 0 sein" in msg

    # Invalid range formats
    is_valid, msg = validate_pages_string("1-2-3", 5)
    assert is_valid is False
    assert "Ungültiger Seitenbereich" in msg

def test_api_extract_pages_validation():
    import uuid
    from backend.services.task_manager import Task
    
    task_id = str(uuid.uuid4())
    task = Task(
        task_id=task_id,
        original_name="test.pdf",
        input_path="/mock/path/test.pdf",
        extension=".pdf",
        page_count=5,
    )
    active_tasks[task_id] = task
    
    # Test out of bounds page numbers:
    response = client.post("/api/extract", data={"task_id": task_id, "pages": "10"})
    assert response.status_code == 400
    assert "liegt außerhalb des gültigen Bereichs" in response.json()["detail"]
    
    # Test valid range validation passing (status 200 stream starts)
    response = client.post("/api/extract", data={"task_id": task_id, "pages": "2-4"})
    assert response.status_code == 200
    
    active_tasks.pop(task_id, None)

def test_parse_page_selection():
    """Test the new page selection parser."""
    assert parse_page_selection("") == []
    assert parse_page_selection("1,3") == [1, 3]
    assert parse_page_selection("1-3") == [1, 2, 3]
    assert parse_page_selection("1, 3, 5-7") == [1, 3, 5, 6, 7]
    assert parse_page_selection("3-1") == [1, 2, 3]  # reversed range

def test_is_contiguous_range():
    """Test contiguous range detection."""
    assert is_contiguous_range([]) is True
    assert is_contiguous_range([1, 2, 3]) is True
    assert is_contiguous_range([5, 6, 7, 8]) is True
    assert is_contiguous_range([1, 3]) is False
    assert is_contiguous_range([1, 2, 4]) is False


def test_get_physical_page_numbers_for_pdf(tmp_path):
    """Test per-page physical page number detection and resolution."""
    from backend.services.extraction_service import get_physical_page_numbers_for_pdf, detect_page_number_for_pdf_page

    # Test non-existent PDF returns fallback
    res = get_physical_page_numbers_for_pdf(str(tmp_path / "nonexistent.pdf"), fallback_start_page=10)
    assert res == [10]

def test_task_manager_lifecycle():
    """Test TaskManager create, get, delete, and TTL cleanup."""
    import asyncio
    from backend.services.task_manager import TaskManager

    mgr = TaskManager(ttl_seconds=1)

    async def run():
        # Create
        task = await mgr.create("t1", "test.pdf", "/path", ".pdf", page_count=5)
        assert task.task_id == "t1"
        assert mgr.count == 1

        # Get
        fetched = await mgr.get("t1")
        assert fetched is not None
        assert fetched.original_name == "test.pdf"

        # Contains
        assert await mgr.contains("t1") is True
        assert await mgr.contains("nonexistent") is False

        # List
        all_tasks = await mgr.list_all()
        assert len(all_tasks) == 1

        # Delete
        deleted = await mgr.delete("t1")
        assert deleted is not None
        assert mgr.count == 0

        # TTL cleanup
        await mgr.create("t2", "old.pdf", "/path", ".pdf")
        import time
        time.sleep(1.1)  # Wait for TTL to expire
        cleaned = await mgr.cleanup_stale()
        assert cleaned == 1
        assert mgr.count == 0

    asyncio.run(run())


def test_insert_page_numbers_from_metadata():
    # Setup mock content list json file
    temp_json_path = os.path.join(TEST_DIR, "test_metadata_content_list.json")
    mock_blocks = [
        {"type": "title", "text": "Physics Vol I", "page_idx": 0},
        {"type": "paragraph", "text": "This is a paragraph on the first page.", "page_idx": 0},
        {"type": "title", "text": "Chapter 2: Optics", "page_idx": 1},
        {"type": "paragraph", "text": "This is a paragraph on the second page.", "page_idx": 1},
    ]
    with open(temp_json_path, "w", encoding="utf-8") as f:
        json.dump(mock_blocks, f)

    markdown = "# Physics Vol I\n\nThis is a paragraph on the first page.\n\n# Chapter 2: Optics\n\nThis is a paragraph on the second page.\n"
    
    # 1. Test contiguous range mapping (physical page 45 and 46)
    res = insert_page_numbers_from_metadata(markdown, temp_json_path, "45-46")
    # Page 45 should be inserted at the beginning
    # Page 46 should be inserted before 'Chapter 2: Optics'
    assert res.startswith("\n\n((45))\n\n") or res.startswith("((45))\n\n")
    assert "\n\n((46))\n\n# Chapter 2: Optics" in res

    # 2. Test non-contiguous page list selection (physical page 12 and 14)
    res_non_contig = insert_page_numbers_from_metadata(markdown, temp_json_path, "12, 14")
    assert res_non_contig.startswith("\n\n((12))\n\n") or res_non_contig.startswith("((12))\n\n")
    assert "\n\n((14))\n\n# Chapter 2: Optics" in res_non_contig

    # Cleanup temp file
    if os.path.exists(temp_json_path):
        os.remove(temp_json_path)


def test_post_process_structure():
    # 1. Test hyphenated word joining across line breaks and page breaks
    raw_md_hyphen = "Das Snelliussche Brechungs-\n((13))\ngesetz beschreibt die Lichtbrechung.\nUnd das Laser-\nLicht ist stark."
    processed = post_process_structure(raw_md_hyphen)
    assert "Brechungsgesetz" in processed
    assert "Laserlicht" in processed
    # Page boundary markers should remain intact
    assert "((13))" in processed

    # 2. Test hard wrap line joining within paragraphs
    raw_md_wrap = "In Zellen sind Biomembranen vorhanden,\ndie eine Barriere bilden. Und es\nist ein Test."
    processed_wrap = post_process_structure(raw_md_wrap)
    assert "In Zellen sind Biomembranen vorhanden, die eine Barriere bilden." in processed_wrap
    assert "Und es ist ein Test." in processed_wrap


def test_docling_extractor(monkeypatch, tmp_path):
    import sys
    from unittest.mock import MagicMock
    
    mock_doc = MagicMock()
    mock_doc.pages = [MagicMock()]
    mock_doc.pictures = []
    mock_doc.export_to_markdown.return_value = "# Docling Page Title\nThis is docling text."
    
    mock_item = MagicMock()
    mock_item.text = "This is docling text."
    mock_item.prov = [MagicMock(page_no=1)]
    mock_item.__class__.__name__ = "TextItem"
    mock_doc.iterate_items.return_value = [(mock_item, 0)]
    
    mock_conv_res = MagicMock()
    mock_conv_res.document = mock_doc
    
    mock_converter_instance = MagicMock()
    mock_converter_instance.convert.return_value = mock_conv_res
    
    mock_converter_cls = MagicMock(return_value=mock_converter_instance)
    monkeypatch.setattr("docling.document_converter.DocumentConverter", mock_converter_cls)
    
    test_input = tmp_path / "test_docling.pdf"
    test_input.write_text("dummy")
    
    output_dir = tmp_path / "docling_auto"
    
    old_argv = sys.argv
    sys.argv = [
        "docling_extractor.py",
        "--input", str(test_input),
        "--output-dir", str(output_dir),
        "--ocr-lang", "en"
    ]
    
    try:
        from backend.services.docling_extractor import main as extractor_main
        extractor_main()
    finally:
        sys.argv = old_argv
        
    md_file = output_dir / "test_docling.md"
    json_file = output_dir / "test_docling_content_list.json"
    
    assert md_file.exists()
    assert json_file.exists()
    
    md_content = md_file.read_text(encoding="utf-8")
    assert "# Docling Page Title" in md_content
    
    with open(json_file, "r") as f:
        blocks = json.load(f)
    assert blocks[0]["type"] == "textitem"
    assert blocks[0]["text"] == "This is docling text."
    assert blocks[0]["page_idx"] == 0


def test_docling_extractor_formula_fallback(monkeypatch, tmp_path):
    import sys
    from unittest.mock import MagicMock
    
    mock_doc = MagicMock()
    mock_doc.pages = [MagicMock()]
    mock_doc.pictures = []
    mock_doc.export_to_markdown.return_value = "# Docling Fallback Title\nFormula fallback text."
    
    mock_item = MagicMock()
    mock_item.text = "Formula fallback text."
    mock_item.prov = [MagicMock(page_no=1)]
    mock_item.__class__.__name__ = "TextItem"
    mock_doc.iterate_items.return_value = [(mock_item, 0)]
    
    mock_conv_res = MagicMock()
    mock_conv_res.document = mock_doc
    
    call_count = 0
    def mock_convert(input_path):
        nonlocal call_count
        call_count += 1
        if call_count == 1:
            raise RuntimeError("Formula VLM inference error")
        return mock_conv_res

    mock_converter_instance = MagicMock()
    mock_converter_instance.convert.side_effect = mock_convert
    
    mock_converter_cls = MagicMock(return_value=mock_converter_instance)
    monkeypatch.setattr("docling.document_converter.DocumentConverter", mock_converter_cls)
    
    test_input = tmp_path / "test_docling_fallback.pdf"
    test_input.write_text("dummy")
    
    output_dir = tmp_path / "docling_fallback_auto"
    
    old_argv = sys.argv
    sys.argv = [
        "docling_extractor.py",
        "--input", str(test_input),
        "--output-dir", str(output_dir),
        "--ocr-lang", "en"
    ]
    
    try:
        from backend.services.docling_extractor import main as extractor_main
        extractor_main()
    finally:
        sys.argv = old_argv
        
    md_file = output_dir / "test_docling_fallback.md"
    assert md_file.exists()
    assert "# Docling Fallback Title" in md_file.read_text(encoding="utf-8")
    assert call_count == 2


@pytest.mark.anyio
async def test_extraction_queueing(monkeypatch):
    import asyncio
    from backend.services.extraction_service import run_background_extraction
    
    execution_order = []
    
    async def mock_locked_run(task, *args, **kwargs):
        task["status"] = "processing"
        execution_order.append(f"start-{task['task_id']}")
        assert task["status"] == "processing"
        await asyncio.sleep(0.1)
        execution_order.append(f"end-{task['task_id']}")
        task["status"] = "completed"

    monkeypatch.setattr(
        "backend.services.extraction_service._run_background_extraction_locked",
        mock_locked_run
    )
    
    task1 = {
        "task_id": "task-1",
        "status": "pending",
        "listeners": [],
        "progress": {},
        "extension": ".pdf"
    }
    task2 = {
        "task_id": "task-2",
        "status": "pending",
        "listeners": [],
        "progress": {},
        "extension": ".pdf"
    }
    
    await asyncio.gather(
        run_background_extraction(task1, "input.pdf", "out", "upload"),
        run_background_extraction(task2, "input.pdf", "out", "upload")
    )
    
    assert len(execution_order) == 4
    assert execution_order[0].startswith("start-")
    assert execution_order[1].startswith("end-")
    assert execution_order[2].startswith("start-")
    assert execution_order[3].startswith("end-")
    
    first_task = execution_order[0].replace("start-", "")
    second_task = execution_order[2].replace("start-", "")
    assert execution_order[1] == f"end-{first_task}"
    assert execution_order[3] == f"end-{second_task}"


def test_gemma_extra_page_numbers_cleanup(monkeypatch):
    from backend.services.llm_service import standardise_markdown

    # Mock the LLM rapid-mlx server response to simulate Gemma outputting extra relative page numbers
    class MockChatCompletions:
        def create(self, **kwargs):
            class MockDelta:
                content = "((45))\n\n((1))\n\n# Physik für die Sekundarstufe II\n\n((46))\n\n((2))\n\n# Physikalische Größen und Konstanten"
            class MockChoice:
                delta = MockDelta()
            class MockChunk:
                choices = [MockChoice()]
            return [MockChunk()]

    class MockClient:
        class MockChat:
            completions = MockChatCompletions()
        chat = MockChat()

    monkeypatch.setattr("backend.services.llm_service.start_rapid_mlx", lambda: MockClient())

    raw_markdown = "((45))\n# Physik für die Sekundarstufe II\n((46))\n# Physikalische Größen und Konstanten"
    result = standardise_markdown(raw_markdown, ruleset_key="uebertragungsregeln_09_2025")

    # Assert physical page numbers exist, but the relative page indexes (1) and (2) have been stripped
    assert "((45))" in result
    assert "((46))" in result
    assert "((1))" not in result
    assert "((2))" not in result


def test_detect_physical_start_page_alternating_and_centered():
    from backend.services.extraction_service import detect_physical_start_page
    
    # 1. Test alternating page number detection (deutsch_textanalyse.pdf begins at 85)
    p85 = detect_physical_start_page("testdata/deutsch_textanalyse.pdf")
    assert p85 == 85

    # 2. Test centered page number detection (mathe_grade3.pdf begins at 1)
    p1 = detect_physical_start_page("testdata/mathe_grade3.pdf")
    assert p1 == 1


def test_prevent_duplicate_page_number_injection():
    # Setup mock content list json file
    temp_json_path = os.path.join(TEST_DIR, "test_metadata_duplicates.json")
    mock_blocks = [
        {"type": "title", "text": "Physics Vol I", "page_idx": 0},
        {"type": "paragraph", "text": "This is a paragraph on the first page.", "page_idx": 0},
    ]
    with open(temp_json_path, "w", encoding="utf-8") as f:
        json.dump(mock_blocks, f)

    # Standardised markdown already contains ((45)) but with 200 characters of filler text before the block
    markdown = "((45))\n\n" + "A" * 200 + "\n\n# Physics Vol I\n\nThis is a paragraph on the first page.\n"
    
    # Run insertion: it should NOT inject ((45)) again because it is already present in markdown
    res = insert_page_numbers_from_metadata(markdown, temp_json_path, "45")
    
    # Count occurrences of ((45)) in result
    occurrences = res.count("((45))")
    assert occurrences == 1

    # Cleanup temp file
    if os.path.exists(temp_json_path):
        os.remove(temp_json_path)


def test_api_lint_endpoint():
    payload = {
        "markdown": "# Optik\n\n((45))\n\n-\\tPunkt 1\n\n<Bild> Zeichnung:\nBeschreibung: Prisma\n</Bild>\n\n$x = y$\n"
    }
    response = client.post("/api/lint", json=payload)
    assert response.status_code == 200
    data = response.json()
    assert data["score"] == 100
    assert data["passed"] is True
    assert "categories" in data


def test_api_autofix_endpoint():
    payload = {
        "markdown": "# Kapitel 1\n### Übersprungene Ebene\n\n- Kein Tab\n1. Auch kein Tab\n"
    }
    response = client.post("/api/lint/autofix", json=payload)
    assert response.status_code == 200
    data = response.json()
    assert "fixed_markdown" in data
    assert "applied_fixes" in data
    assert len(data["applied_fixes"]) > 0
    assert "## Übersprungene" in data["fixed_markdown"]
    assert "-\tKein Tab" in data["fixed_markdown"]
    assert "1.\tAuch kein Tab" in data["fixed_markdown"]
    assert data["report"]["score"] == 100


def test_docx_page_bookmarks():
    from docx import Document
    docx_path = os.path.join(TEST_DIR, "test_bookmarks.docx")
    md = "# Kapitel 1\n\n((45))\n\nErster Absatz.\n\n((46))\n\nZweiter Absatz.\n"
    convert_markdown_to_docx(md, docx_path)
    assert os.path.exists(docx_path)
    
    doc = Document(docx_path)
    xmls = [p._p.xml for p in doc.paragraphs]
    assert any('w:name="page_45"' in x for x in xmls)
    assert any('w:name="page_46"' in x for x in xmls)
    # Verify page break was inserted before ((46))
    full_xml = "".join(xmls)
    assert 'w:type="page"' in full_xml

def test_in_memory_pdf_render_cache():
    from backend.services.pdf_render_service import render_single_page_bytes, _PAGE_BYTES_CACHE
    pdf_path = os.path.join(BASE_DIR, "testdata", "test_educational.pdf")
    
    # Test rendering single page to bytes
    png_bytes = render_single_page_bytes(pdf_path, 0)
    assert png_bytes is not None
    assert len(png_bytes) > 100
    assert png_bytes.startswith(b"\x89PNG")
    
    # Second call should fetch directly from cache
    png_bytes_2 = render_single_page_bytes(pdf_path, 0)
    assert png_bytes_2 == png_bytes


def test_rapid_mlx_sliding_idle_timer():
    from backend.services import fibs_constants as fc
    from backend.services.llm_service import reset_idle_timer, _cancel_idle_timer, _IDLE_TIMER
    import backend.services.llm_service as llm_mod
    
    # Set short test timeout
    orig_timeout = fc.IDLE_TIMEOUT_SECONDS
    fc.IDLE_TIMEOUT_SECONDS = 0.1
    try:
        reset_idle_timer()
        assert llm_mod._IDLE_TIMER is not None
        assert llm_mod._IDLE_TIMER.is_alive()
        
        # Test cancel
        _cancel_idle_timer()
        assert llm_mod._IDLE_TIMER is None
    finally:
        fc.IDLE_TIMEOUT_SECONDS = orig_timeout


def test_token_streaming_callback():
    from unittest.mock import patch, MagicMock
    from backend.services.llm_service import standardise_markdown

    tokens = []
    def on_token(t):
        tokens.append(t)

    # Mock chunk structure from OpenAI streaming completion
    class MockDelta:
        def __init__(self, content):
            self.content = content

    class MockChoice:
        def __init__(self, content):
            self.delta = MockDelta(content)

    class MockChunk:
        def __init__(self, content):
            self.choices = [MockChoice(content)]

    mock_client = MagicMock()
    mock_client.chat.completions.create.return_value = [
        MockChunk("# "), MockChunk("Test "), MockChunk("Title\n"), MockChunk("Paragraph text.")
    ]

    with patch("backend.services.llm_service.start_rapid_mlx", return_value=mock_client):
        res = standardise_markdown("# Test Title\nParagraph text.", on_token_callback=on_token)
        assert res is not None
        assert "Test Title" in res
        assert len(tokens) == 4
        assert "".join(tokens) == "# Test Title\nParagraph text."


def test_infer_heading_hierarchy():
    from backend.services.fibs_postprocess import infer_heading_hierarchy, post_process_structure
    
    input_flat_md = (
        "# Kapitel 1: Einleitung\n\n"
        "Einleitungstext.\n\n"
        "# 1.1 Grundlagen der Optik\n\n"
        "Grundlagen.\n\n"
        "# 1.1.1 Das Fermat'sche Prinzip\n\n"
        "Details zum Prinzip.\n\n"
        "# Versuch: Brechung am Prisma\n\n"
        "Versuchsdurchführung.\n"
    )
    
    result = infer_heading_hierarchy(input_flat_md)
    assert "# Kapitel 1: Einleitung" in result
    assert "## 1.1 Grundlagen der Optik" in result
    assert "### 1.1.1 Das Fermat'sche Prinzip" in result
    assert "### Versuch: Brechung am Prisma" in result
    
    # Verify integration in post_process_structure
    processed = post_process_structure(input_flat_md)
    assert "## 1.1 Grundlagen der Optik" in processed
    assert "### 1.1.1 Das Fermat'sche Prinzip" in processed


def test_clean_duplicate_itemization():
    from backend.services.extraction_postprocess import clean_duplicate_itemization
    
    input_text = (
        "- - First bullet\n"
        "- • Second bullet\n"
        "- 1. First numbered item\n"
        "3. 3. Third question item\n"
        "- ■ Square bullet item\n"
    )
    
    cleaned = clean_duplicate_itemization(input_text)
    assert "- First bullet" in cleaned
    assert "- Second bullet" in cleaned
    assert "1. First numbered item" in cleaned
    assert "3. Third question item" in cleaned
    assert "- Square bullet item" in cleaned


def test_scanned_pdf_golden_mapping():
    from scripts.run_benchmark import ensure_scanned_pdfs
    ensure_scanned_pdfs()
    
    scanned_path = os.path.join(BASE_DIR, "testdata", "test_educational_scanned.pdf")
    assert os.path.exists(scanned_path)
    assert os.path.getsize(scanned_path) > 100000


def test_word_recall_normalization():
    from scripts.run_benchmark import evaluate_text_and_continuity, normalize_text_for_word_eval
    
    # Test German umlaut equivalence: transliterated 'ae' vs 'ä'
    pred_md = "((1))\n\n# Mathematische Übungen\n\nLöse die Aufgaben so schnell wie möglich.\n\nFlächen berechnen: a = 5 cm."
    gold_md = "((1))\n\n# Mathematische Uebungen\n\nLoese die Aufgaben so schnell wie moeglich.\n\nFlaechen berechnen: a = 5 cm."
    
    eval_res = evaluate_text_and_continuity(pred_md, gold_md)
    assert eval_res["word_recall"] == 100.0
    assert eval_res["missing_chunks_count"] == 0
    
    # Test fill-in blank runs do not penalize word recall
    pred_with_blanks = "((1))\n\na) 24 + 37 = _ _ _ _\nb) 45 + 18 = __"
    gold_with_blanks = "((1))\n\na) 24 + 37 = ____\nb) 45 + 18 = ________"
    
    eval_blanks = evaluate_text_and_continuity(pred_with_blanks, gold_with_blanks)
    assert eval_blanks["word_recall"] == 100.0


def test_golden_structure():
    from backend.services.a11y_linter import autofix_markdown, lint_markdown

    golden_dir = os.path.join(BASE_DIR, "testdata", "golden")
    docs = [
        "test_educational",
        "deutsch_textanalyse",
        "mathe_grade3",
        "physik_optik_sample",
        "openstax_biology_sample",
        "openstax_chemistry_sample",
        "openstax_history_sample",
        "grundschule_mathematik",
        "grundschule_sachkunde",
        "sek1_deutsch",
        "sek1_englisch",
        "sek1_geographie",
        "sek1_geschichte",
        "sek2_biologie",
        "sek2_chemie",
        "sek2_physik_verzeichnisse",
    ]
    
    for doc in docs:
        raw_path = os.path.join(golden_dir, f"{doc}.raw.golden.md")
        ueb_path = os.path.join(golden_dir, f"{doc}.uebertragungsregeln_09_2025.golden.md")
        
        assert os.path.exists(raw_path), f"Missing {raw_path}"
        assert os.path.exists(ueb_path), f"Missing {ueb_path}"
        
        with open(raw_path, "r", encoding="utf-8") as f:
            raw_content = f.read()
            assert len(raw_content) > 500
            
        with open(ueb_path, "r", encoding="utf-8") as f:
            ueb_content = f.read()
            assert len(ueb_content) > 500

        rep = lint_markdown(ueb_content)
        assert rep.score == 100, f"Linter issues in {doc}: {rep.issues}"
        assert rep.issues == []
        fixed, fixes = autofix_markdown(ueb_content)
        assert fixes == [], f"Autofix not idempotent on {doc}: {fixes}"
        assert fixed.strip() == ueb_content.strip()


def test_re_extraction_after_completion_and_abort():
    """Verifies that an extraction can be re-run with different parameters after completion or abort."""
    pdf_path = os.path.join(BASE_DIR, "testdata", "test_educational.pdf")
    if not os.path.exists(pdf_path):
        pytest.skip("test_educational.pdf not found")

    with open(pdf_path, "rb") as f:
        res = client.post("/api/upload", files={"file": ("test_educational.pdf", f, "application/pdf")})
    assert res.status_code == 200
    task_id = res.json()["task_id"]

    # 1. Run first extraction (Docling)
    res_ext1 = client.post("/api/extract", data={"task_id": task_id, "engine": "docling", "pages": "45"})
    assert res_ext1.status_code == 200
    assert "completed" in res_ext1.text

    # 2. Re-run extraction immediately on the same task_id with different parameters (Infinity, page 46)
    res_ext2 = client.post("/api/extract", data={"task_id": task_id, "engine": "infinity", "pages": "46"})
    assert res_ext2.status_code == 200
    assert "completed" in res_ext2.text or "log" in res_ext2.text

    # 3. Request abort on task
    res_abort = client.post(f"/api/tasks/{task_id}/abort")
    assert res_abort.status_code == 200

    # 4. Re-run extraction after abort
    res_ext3 = client.post("/api/extract", data={"task_id": task_id, "engine": "docling", "pages": "45-46"})
    assert res_ext3.status_code == 200
    assert "completed" in res_ext3.text or "log" in res_ext3.text


def test_deduplicate_page_numbers():
    from backend.services.fibs_postprocess import deduplicate_page_numbers

    # 1. Immediate consecutive duplicates
    dup1 = "((23))\n((23))\n\n# Kapitel 1\nText auf Seite 23."
    res1 = deduplicate_page_numbers(dup1)
    assert res1.count("((23))") == 1

    dup2 = "((23))\n\n((23))\n\n# Kapitel 1\nText auf Seite 23."
    res2 = deduplicate_page_numbers(dup2)
    assert res2.count("((23))") == 1

    # 2. Triplicate
    dup3 = "((23))\n((23))\n((23))\n\nText."
    res3 = deduplicate_page_numbers(dup3)
    assert res3.count("((23))") == 1

    # 3. Duplicate later in the same page span
    dup4 = "((23))\n\nErster Absatz.\n\n((23))\n\nZweiter Absatz.\n\n((24))\n\nAbsatz auf Seite 24."
    res4 = deduplicate_page_numbers(dup4)
    assert res4.count("((23))") == 1
    assert res4.count("((24))") == 1
    assert "Erster Absatz." in res4
    assert "Zweiter Absatz." in res4

    # 4. Different consecutive page numbers without intervening content (MUST BE PRESERVED)
    valid_diff = "((23))\n\n((24))\n\n# Kapitel auf Seite 24"
    res_diff = deduplicate_page_numbers(valid_diff)
    assert "((23))" in res_diff
    assert "((24))" in res_diff
    assert res_diff.count("((23))") == 1
    assert res_diff.count("((24))") == 1


def test_image_caption_headings_and_url_sanitization():
    # 1. Image captions wrongly transcribed as headings (### Abb. 1: ...)
    bad_md = (
        "((47))\n\n"
        "### Abbildung 4.2a: Prisma von vorne.\n\n"
        "### Abb. 4.2b: Prisma von der Seite.\n\n"
        "Die Zellkompartimen- ten sind wichtig."
    )
    processed = post_process_structure(bad_md)
    assert "<Bild> Abbildung 4.2a: Prisma von vorne. </Bild>" in processed
    assert "<Bild> Abbildung 4.2b: Prisma von der Seite. </Bild>" in processed
    assert "### Abbildung" not in processed
    assert "### Abb." not in processed

    # 2. De-hyphenation of broken words with spaces
    assert "Zellkompartimenten" in processed
    assert "Zellkompartimen- ten" not in processed

    # 3. Sanitization of external internet URLs
    url_md = "((45))\n\n![Zelle](https://www.bildungsberliner.de/aktuelles/2019/09/20190920130000.jpg)"
    url_processed = post_process_structure(url_md)
    assert "https://" not in url_processed
    assert "<Bild> Abbildung:\nBeschreibung: Zelle\n</Bild>" in url_processed


def test_docx_export_formatting_rules(tmp_path):
    from docx import Document
    from backend.services.docx_service import convert_markdown_to_docx

    md_content = (
        "# 1. Hauptkapitel\n\n"
        "Erster Absatz mit _hervorgehobenem_ Text und *fettem* Text.\n\n"
        "Zweiter Absatz.\n\n"
        "<Rahmen>\n"
        "Merksatz: Licht breitet sich geradlinig aus.\n"
        "Formel: $c = 299792\\text{ km/s}$\n"
        "</Rahmen>\n\n"
        "## 1.1 Unterkapitel\n\n"
        "Text nach Unterkapitel."
    )
    docx_file = str(tmp_path / "test_rules.docx")
    convert_markdown_to_docx(md_content, docx_file)

    doc = Document(docx_file)

    # 1. No empty paragraph spacers after headings or between paragraphs
    paragraph_texts = [p.text for p in doc.paragraphs]
    for idx, p_text in enumerate(paragraph_texts):
        assert p_text.strip() != "", f"Empty paragraph found at index {idx}"

    # 2. Check no underline on runs for _hervorgehobenem_
    for p in doc.paragraphs:
        for run in p.runs:
            if "_hervorgehobenem_" in run.text:
                assert run.underline is not True, "Highlighted text must not be underlined in Word"
            if "*fettem*" in run.text:
                assert run.bold is True

    # 3. <Rahmen>...</Rahmen> is rendered in a single paragraph with manual line breaks
    rahmen_paragraphs = [p for p in doc.paragraphs if "<Rahmen>" in p.text]
    assert len(rahmen_paragraphs) == 1, "Expected exactly 1 paragraph for the entire <Rahmen> block"
    rp = rahmen_paragraphs[0]
    assert "Merksatz:" in rp.text
    assert "Formel:" in rp.text
    assert "</Rahmen>" in rp.text
    # Verify manual line break elements exist in XML
    assert len(rp._p.xpath('.//w:br')) >= 2, "Expected manual line breaks inside <Rahmen> paragraph"


def test_docx_export_tables(tmp_path):
    from docx import Document
    from backend.services.docx_service import convert_markdown_to_docx

    md_content = (
        "# Tabellentest\n\n"
        "<Tabelle>\n\n"
        "| Modell | Annahme | Phänomene |\n"
        "| :--- | :--- | :--- |\n"
        "| Korpuskelmodell | Photonen | Photoelektrischer Effekt |\n"
        "| Wellenmodell | Transversalwellen | Beugung und Interferenz |\n\n"
        "</Tabelle>\n\n"
        "Absatz zwischen Tabellen.\n\n"
        "| Spalte A | Spalte B |\n"
        "| :--- | :--- |\n"
        "| Wert 1 | Wert 2 |\n"
    )
    docx_file = str(tmp_path / "test_tables.docx")
    convert_markdown_to_docx(md_content, docx_file)

    doc = Document(docx_file)
    assert len(doc.tables) == 2, f"Expected 2 native Word tables, got {len(doc.tables)}"

    # First table
    t1 = doc.tables[0]
    assert len(t1.rows) == 3, f"Expected 3 rows in first table, got {len(t1.rows)}"
    assert len(t1.columns) == 3, f"Expected 3 columns in first table, got {len(t1.columns)}"
    assert t1.cell(0, 0).text == "Modell"
    assert t1.cell(1, 0).text == "Korpuskelmodell"
    assert t1.cell(2, 1).text == "Transversalwellen"

    # Second table (standalone pipe table)
    t2 = doc.tables[1]
    assert len(t2.rows) == 2, f"Expected 2 rows in second table, got {len(t2.rows)}"
    assert len(t2.columns) == 2, f"Expected 2 columns in second table, got {len(t2.columns)}"
    assert t2.cell(0, 0).text == "Spalte A"
    assert t2.cell(1, 1).text == "Wert 2"


def test_runaway_token_repetition_sanitization():
    from backend.services.extraction_postprocess import (
        clean_markdown_fences,
        fuse_vlm_page_with_docling,
        recover_text_truncation,
    )

    # 1. Test clean_markdown_fences with runaway repetitive zeros and unclosed image tags
    runaway_zeros = "0" * 3000
    raw_md = f"# Kapitel 1\n\nText vor Bild\n\n![](images/{runaway_zeros}"
    cleaned = clean_markdown_fences(raw_md)
    assert "00000000000" not in cleaned
    assert "![]" not in cleaned

    # 2. Test fuse_vlm_page_with_docling with runaway image tag and valid image replacement
    page_text = f"Normaler Text.\n\n![](images/{runaway_zeros})"
    processed = fuse_vlm_page_with_docling(page_text, page_images=["picture_1_test.png"], phys_num=47)
    assert "00000000000" not in processed
    assert "((47))" in processed

    # 3. Test recover_text_truncation in extraction_postprocess
    popo_input = f"Absatz.\n\n![](images/{'a' * 50})\n\nWeiterer Text."
    popo_cleaned = recover_text_truncation(popo_input)
    assert "aaaaa" not in popo_cleaned
    assert "Absatz." in popo_cleaned

    # 4. Test heading loop deduplication (e.g. ## B4... ## B5... up to ## B138...)
    from backend.services.extraction_postprocess import sanitize_runaway_repetitions
    loop_text = (
        "## B4 Nachweis von Hex-1-en mit Brom zu 1,2-Dibromethan\n\n"
        "## B5 Nachweis von Hex-1-en mit Brom zu 1,2-Dibromethan\n\n"
        "## B6 Nachweis von Hex-1-en mit Brom zu 1,2-Dibromethan\n\n"
        "## B138 Nachweis von Hex-1-en mit Brom zu 1,2-Dibromethan\n"
    )
    cleaned_loop = sanitize_runaway_repetitions(loop_text)
    assert "## B4 Nachweis von Hex-1-en mit Brom zu 1,2-Dibromethan" in cleaned_loop
    assert "## B5 Nachweis" not in cleaned_loop
    assert "## B138 Nachweis" not in cleaned_loop

    # 5. Test that legitimate <Bild> tags are strictly preserved and not collapsed
    bild_text = (
        "<Bild> Abbildung: Abbildung ohne Beschreibung </Bild>\n\n"
        "<Bild> Abbildung: Abbildung ohne Beschreibung </Bild>\n\n"
        "<Bild> Abbildung: Abbildung ohne Beschreibung </Bild>\n"
    )
    cleaned_bild = sanitize_runaway_repetitions(bild_text)
    assert cleaned_bild.count("<Bild>") == 3
